#!/usr/bin/env python3
"""slides_content.json + テンプレートpptx から 日本語版/英語版のスライドを生成する。

内容の判断（何を書くか）はこのスクリプトの責務ではない。SKILL.md の指示に従って
Claude が作成した slides_content.json を、テンプレートのレイアウトへ機械的に
流し込むだけ。

使い方:
    python build_pptx.py --content <weekly_dir>/slides_content.json --outdir <weekly_dir>
    (--template を省略すると templates/ 配下の唯一の .pptx を自動使用する)
"""
import argparse
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu
from PIL import Image

SCRIPT_DIR = Path(__file__).resolve().parent
TEMPLATES_DIR = SCRIPT_DIR.parent / "templates"

# テンプレートの表紙は「サンプルスライドの実データ」としてこの文言を持っていた
# （レイアウト自体は空のプレースホルダしか持たない）。サンプルスライドは
# clear_slides() で削除するため、既定値としてここに保持しておく。
# 所属が変わったら書き換えること。
DEFAULT_AFFILIATION = {
    "ja": "1 大阪大学大学院情報科学研究科\n2 理化学研究所計算科学研究センター",
    "en": "1 Graduate School of Information Science and Technology, Osaka University\n"
          "2 RIKEN Center for Computational Science",
}


def find_template() -> Path:
    candidates = sorted(TEMPLATES_DIR.glob("*.pptx"))
    if not candidates:
        sys.exit(f"テンプレートが見つかりません: {TEMPLATES_DIR} に .pptx を配置してください")
    if len(candidates) > 1:
        sys.exit(
            "テンプレートが複数あります。--template で明示してください: "
            + ", ".join(str(c) for c in candidates)
        )
    return candidates[0]


def clear_slides(prs):
    """テンプレート同梱のサンプルスライドを全て削除する。

    sldIdLst から参照を外すだけだと元のスライドパート（slide1.xml等）が
    パッケージに残ったままになり、新規追加スライドとパート名が衝突して
    不正なpptxになる。drop_rel で関連パートごと削除する。
    """
    xml_slides = prs.slides._sldIdLst
    for sld in list(xml_slides):
        prs.part.drop_rel(sld.rId)
        xml_slides.remove(sld)


def placeholder_by_idx(slide, idx):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            return ph
    return None


def set_first_run_text(shape, text):
    """既存の1つ目のrunのテキストのみ差し替え、フォント書式（bold/sizeなど）を維持する。"""
    tf = shape.text_frame
    p = tf.paragraphs[0]
    if not p.runs:
        p.add_run()
    p.runs[0].text = text
    # 同じ段落に他のrunが残っていれば（例: 表紙の名前欄の上付き番号）空にする
    for r in p.runs[1:]:
        r.text = ""


def set_multiline_text(shape, text):
    """改行区切りのテキストを段落ごとに設定する（所属欄など）。"""
    tf = shape.text_frame
    tf.clear()
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line


def set_bullets(shape, bullets, lang):
    tf = shape.text_frame
    tf.clear()
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = b.get("level", 0)
        run = p.add_run()
        run.text = b["text"][lang]


def fit_box(img_path, left, top, width, height, pad_ratio=0.04):
    with Image.open(img_path) as im:
        iw, ih = im.size
    pad = int(min(width, height) * pad_ratio)
    box_w = width - 2 * pad
    box_h = height - 2 * pad
    scale = min(box_w / iw, box_h / ih)
    w = int(iw * scale)
    h = int(ih * scale)
    l = left + pad + (box_w - w) // 2
    t = top + pad + (box_h - h) // 2
    return l, t, w, h


def add_cover_slide(prs, layout, cover, lang):
    slide = prs.slides.add_slide(layout)

    if "year" not in cover:
        sys.exit("cover.year が指定されていません（例: {\"ja\": \"2026\", \"en\": \"2026\"}）")

    single_run_fields = {
        0: cover.get("title"),
        1: cover.get("event_name"),
        2: cover.get("year"),
        4: cover.get("authors"),
        5: cover.get("date"),
    }
    for idx, value in single_run_fields.items():
        if value is None:
            continue
        ph = placeholder_by_idx(slide, idx)
        if ph is None:
            continue
        set_first_run_text(ph, value[lang])

    affiliation = cover.get("affiliation", DEFAULT_AFFILIATION)
    aff_ph = placeholder_by_idx(slide, 3)
    if aff_ph is not None:
        set_multiline_text(aff_ph, affiliation[lang])
    return slide


def add_content_slide(prs, layout_map, section, lang, base_dir):
    layout_name = section.get("layout", "OBJECT")
    if layout_name not in layout_map:
        sys.exit(f"section '{section.get('id')}': レイアウト '{layout_name}' が見つかりません"
                  f"（利用可能: {', '.join(layout_map)}）")
    slide = prs.slides.add_slide(layout_map[layout_name])

    title_ph = placeholder_by_idx(slide, 0)
    if title_ph is not None and "title" in section:
        set_first_run_text(title_ph, section["title"][lang])

    body_ph = placeholder_by_idx(slide, 1)
    if body_ph is not None and section.get("bullets"):
        set_bullets(body_ph, section["bullets"], lang)

    images = section.get("images", [])
    if images:
        if layout_name != "2カラム":
            sys.exit(f"section '{section.get('id')}': images は layout='2カラム' でのみ対応")
        img_ph = placeholder_by_idx(slide, 2)
        box_left, box_top = img_ph.left, img_ph.top
        box_w, box_h = img_ph.width, img_ph.height
        slot_h = box_h // len(images)
        for i, img in enumerate(images):
            path = base_dir / img["path"]
            if not path.exists():
                sys.exit(f"画像が見つかりません: {path}")
            slot_top = box_top + i * slot_h
            caption = img.get("caption")
            cap_h = Emu(int(0.35 * 914400)) if caption else 0
            l, t, w, h = fit_box(path, box_left, slot_top, box_w, slot_h - cap_h)
            slide.shapes.add_picture(str(path), l, t, width=w, height=h)
            if caption:
                tb = slide.shapes.add_textbox(box_left, slot_top + (slot_h - cap_h), box_w, cap_h)
                tf = tb.text_frame
                tf.word_wrap = True
                run = tf.paragraphs[0].add_run()
                run.text = caption[lang]
                run.font.size = Emu(int(0.14 * 914400))
    return slide


def build_one(template_path, content, lang, out_path, base_dir):
    prs = Presentation(str(template_path))
    layout_map = {l.name: l for l in prs.slide_layouts}
    clear_slides(prs)

    add_cover_slide(prs, layout_map["表紙"], content["cover"], lang)
    for section in content["sections"]:
        add_content_slide(prs, layout_map, section, lang, base_dir)

    prs.save(str(out_path))
    print(f"generated: {out_path}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--content", required=True, type=Path, help="slides_content.json")
    ap.add_argument("--outdir", required=True, type=Path)
    ap.add_argument("--template", type=Path, default=None)
    ap.add_argument("--base-dir", type=Path, default=Path("."),
                     help="images内の相対パスの基準ディレクトリ（デフォルト: カレントディレクトリ＝リポジトリルート）")
    args = ap.parse_args()

    template_path = args.template or find_template()
    content = json.loads(args.content.read_text(encoding="utf-8"))
    args.outdir.mkdir(parents=True, exist_ok=True)

    build_one(template_path, content, "ja", args.outdir / "slides_ja.pptx", args.base_dir)
    build_one(template_path, content, "en", args.outdir / "slides_en.pptx", args.base_dir)


if __name__ == "__main__":
    main()
